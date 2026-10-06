"""Trust, document scope and continuation regressions for the personal product."""

from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

import pytest

from researchpilot.agents.base import build_runtime
from researchpilot.agents.researcher import ResearchAgent
from researchpilot.agents.support import direct_support, quote_grounded, support_check
from researchpilot.agents.verifier import VerifierAgent
from researchpilot.agents.writer import WriterAgent
from researchpilot.llm.mock_provider import MockLLMProvider
from researchpilot.pipeline import ResearchPipeline, ResearchScopeError
from researchpilot.rag.loader import DocumentLoader
from researchpilot.schemas import (
    CitationCheck,
    Evidence,
    EvidenceBundle,
    FinalReport,
    ReportClaim,
    ReportSection,
    ResearchPlan,
    ResearchRequest,
    Subtask,
    VerificationReport,
)


def test_source_matching_preserves_nfkc_punctuation_without_losing_conditions():
    source = "仅当管理员批准时,本产品不支持离线部署。"
    quote = "仅当管理员批准时，本产品不支持离线部署。"
    assert quote_grounded(quote, source)
    assert direct_support(quote, source)
    assert not direct_support("本产品不支持离线部署。", source)
    assert not quote_grounded("本产品支持离线部署。", source)


@pytest.mark.parametrize(
    "claim,quote",
    [
        ("本产品支持离线部署，无须连接云端服务。", "本产品不支持离线部署，必须连接云端服务。"),
        ("本产品不支持离线部署。", "本产品支持离线部署。"),
        ("The service supports offline deployment.", "The service does not support offline deployment."),
        ("本产品收费 100 元。", "本产品收费 10 元。"),
        ("本产品将在 2026 年发布。", "本产品将在 2025 年发布。"),
    ],
)
def test_obvious_reversals_cannot_be_supported_even_by_a_positive_model(claim, quote):
    model = CitationCheck(evidence_id="E1", statement=claim, status="supported")
    assert support_check("E1", claim, quote, model).status == "unsupported"


def test_paraphrases_require_semantics_and_consume_the_model_verdict():
    claim, quote = "该服务离线可用。", "无需网络也能使用该服务。"
    assert not direct_support(claim, quote)
    assert support_check("E1", claim, quote).status == "weak"
    model = CitationCheck(evidence_id="E1", statement=claim, status="supported", reason="same meaning")
    assert support_check("E1", claim, quote, model).status == "supported"
    model.status = "unsupported"
    assert support_check("E1", claim, quote, model).status == "unsupported"
    model.statement = "另一条结论"
    assert support_check("E1", claim, quote, model).status == "weak"


def test_source_grounding_catches_negation_removed_from_quote():
    assert not quote_grounded("支持离线部署", "该产品不支持离线部署。")
    assert not quote_grounded("收费 100 元", "该产品收费 10 元。")
    assert quote_grounded("该产品不支持离线部署。", "说明：该产品不支持离线部署。")


def test_temporal_roles_and_conflicting_sources_need_a_semantic_verdict():
    quote = "2025 年收入 10 万，2026 年收入 20 万。"
    claim = "2025 年收入 20 万，2026 年收入 10 万。"
    rejected = CitationCheck(evidence_id="E1", statement=claim, status="unsupported", reason="years swapped")
    assert support_check("E1", claim, quote, rejected).status == "unsupported"
    quotes = "新版本支持离线部署。\n旧版本不支持离线部署。"
    claim = "所有版本都支持离线部署。"
    weak = CitationCheck(evidence_id="E1", statement=claim, status="weak", reason="versions disagree")
    assert support_check("E1", claim, quotes, weak).status != "supported"


@pytest.mark.parametrize(
    "claim,quote",
    [
        ("本产品支持离线部署。", "仅当管理员批准时，本产品支持离线部署。"),
        (
            "Offline deployment is supported.",
            "If approved by the administrator, offline deployment is supported.",
        ),
    ],
)
def test_conditions_cannot_be_removed_by_extracting_a_comma_clause(claim, quote):
    assert not direct_support(claim, quote)
    assert support_check("E1", claim, quote).status == "weak"
    assert support_check("E1", claim, claim, source_context=quote).status == "weak"
    assert support_check("E1", quote, quote, source_context=quote).status == "supported"


def _plan(question="本产品是否支持离线部署？"):
    return ResearchPlan(
        objective=question,
        subtasks=[
            Subtask(
                id="S1",
                question=question,
                intent="knowledge_search",
                tools=["knowledge_search"],
                expected_output="证据",
            )
        ],
        max_iterations=1,
    )


def test_verifier_skips_model_for_direct_extracts_and_rejects_fabricated_quotes(
    settings, knowledge_base, monkeypatch
):
    runtime = build_runtime(
        task_id="trust", question="离线部署", settings=settings, knowledge_base=knowledge_base
    )
    quote = "本产品不支持离线部署，必须连接云端服务。"
    runtime.sources.register_item(
        {"source_id": "local", "title": "产品说明", "content": quote}, tool="knowledge_search"
    )

    def no_call(*args, **kwargs):
        pytest.fail("direct extracts should not spend a model call")

    monkeypatch.setattr(runtime.llm, "run", no_call)
    bundle = EvidenceBundle(
        evidence=[Evidence(id="E1", subtask_id="S1", claim=quote, quote=quote, source_id="local")]
    )
    assert VerifierAgent(runtime).run(_plan(), bundle).checks[0].status == "supported"
    bundle.evidence[0].quote = "本产品支持离线部署，必须连接云端服务。"
    bundle.evidence[0].claim = bundle.evidence[0].quote
    assert VerifierAgent(runtime).run(_plan(), bundle).checks[0].status == "unsupported"


def test_verifier_sends_source_conditions_even_when_claim_equals_quote(settings, knowledge_base, monkeypatch):
    runtime = build_runtime(
        task_id="conditions", question="离线部署", settings=settings, knowledge_base=knowledge_base
    )
    claim, source = "本产品支持离线部署。", "仅当管理员批准时，本产品支持离线部署。"
    runtime.sources.register_item({"source_id": "local", "content": source}, tool="knowledge_search")
    bundle = EvidenceBundle(evidence=[Evidence(id="E1", claim=claim, quote=claim, source_id="local")])
    captured = []

    def review(*args, **kwargs):
        captured.append(kwargs["user"])
        return SimpleNamespace(
            value=VerificationReport(
                checks=[
                    CitationCheck(
                        evidence_id="E1", statement=claim, status="unsupported", reason="condition omitted"
                    )
                ]
            )
        )

    monkeypatch.setattr(runtime.llm, "run", review)
    assert VerifierAgent(runtime).run(_plan(), bundle).checks[0].status == "unsupported"
    assert len(captured) == 1 and source in captured[0] and "source_context" in captured[0]
    monkeypatch.setattr(runtime.llm, "run", lambda *a, **k: SimpleNamespace(value=VerificationReport()))
    assert VerifierAgent(runtime).run(_plan(), bundle).checks[0].status == "weak"


@pytest.mark.parametrize("conditional", [False, True])
def test_final_report_cannot_reverse_verified_evidence(settings, knowledge_base, conditional):
    runtime = build_runtime(
        task_id="final-check", question="离线部署", settings=settings, knowledge_base=knowledge_base
    )
    quote = "本产品不支持离线部署，必须连接云端服务。"
    evidence = Evidence(id="E1", claim=quote, quote=quote, source_id="local")
    invented = "本产品支持离线部署，无须连接云端服务。"
    if conditional:
        quote, invented = "仅当管理员批准时，本产品支持离线部署。", "本产品支持离线部署。"
        evidence = evidence.model_copy(update={"claim": quote, "quote": quote})
    report = FinalReport(
        title="产品资料",
        executive_summary=invented,
        conclusions=[ReportClaim(statement=invented, evidence_ids=["E1"])],
        sections=[ReportSection(heading="部署", body=invented, evidence_ids=["E1"])],
    )
    checked = WriterAgent(runtime)._audit_report(report, [evidence], _plan())
    assert all(
        (invented != text if conditional else invented not in text)
        for text in [checked.executive_summary, checked.sections[0].body, checked.conclusions[0].statement]
    )
    assert quote in checked.executive_summary
    assert checked.rejected_claims
    assert all(check.status == "supported" for check in checked.support_checks)


def test_missing_final_review_falls_back_to_source_text(settings, knowledge_base, monkeypatch):
    runtime = build_runtime(
        task_id="missing-review", question="离线部署", settings=settings, knowledge_base=knowledge_base
    )
    runtime.llm.provider.name = "semantic-test"
    monkeypatch.setattr(runtime.llm, "run", lambda *a, **k: SimpleNamespace(value=VerificationReport()))
    evidence = Evidence(id="E1", claim="无需联网。", quote="无需联网。", source_id="local")
    report = FinalReport(
        title="资料",
        executive_summary="全球最好的工具。",
        conclusions=[ReportClaim(statement="这是全球最好的工具。", evidence_ids=["E1"])],
    )
    checked = WriterAgent(runtime)._audit_report(report, [evidence], _plan())
    assert checked.conclusions[0].statement == "无需联网。"
    assert "全球最好" not in checked.executive_summary


@pytest.mark.parametrize("field", ["conclusion", "section", "limitation", "recommendation"])
def test_missing_units_are_rejected_in_every_report_field(settings, knowledge_base, monkeypatch, field):
    runtime = build_runtime(
        task_id="units", question="收入", settings=settings, knowledge_base=knowledge_base
    )
    runtime.llm.provider.name = "semantic-test"
    quote = "2025 年收入为 10 万元，2026 年收入为 20 万元。"
    runtime.sources.register_item({"source_id": "local", "content": quote}, tool="document_reader")
    # The extracted quote omits the unit; the complete source must still be used.
    evidence = Evidence(id="E1", claim="收入记录", quote="收入记录", source_id="local")
    invented = "资料未提供收入单位。"
    report = FinalReport(title="报告", executive_summary="")
    if field == "conclusion":
        report.conclusions = [ReportClaim(statement=invented, evidence_ids=["E1"])]
    elif field == "section":
        report.sections = [ReportSection(heading="收入", body=invented, evidence_ids=["E1"])]
    elif field == "limitation":
        report.limitations = [invented]
    else:
        report.recommendations = ["由于资料未提供收入单位，建议补充收入单位。"]
    captured = []

    def approve(*args, **kwargs):
        items = kwargs["hints"]["evidence"]
        captured.extend(items)
        return SimpleNamespace(
            value=VerificationReport(
                checks=[
                    CitationCheck(evidence_id=item["id"], statement=item["claim"], status="supported")
                    for item in items
                ]
            )
        )

    monkeypatch.setattr(runtime.llm, "run", approve)
    checked = WriterAgent(runtime)._audit_report(report, [evidence], _plan("收入"))
    displayed = checked.model_dump(exclude={"rejected_claims", "support_checks"})
    assert "未提供收入单位" not in str(displayed)
    assert any(quote == item["source_context"] for item in captured)
    assert checked.rejected_claims


def test_revenue_before_launch_is_not_a_proven_contradiction():
    quote = "2025 年收入为 10 万元。产品于 2026 年 9 月 10 日上线。"
    claim = "2025 年收入与 2026 年 9 月 10 日上线存在时间矛盾。"
    approved = CitationCheck(evidence_id="C0", statement=claim, status="supported")
    assert support_check("C0", claim, quote, approved).status == "unsupported"


@pytest.mark.parametrize(
    "claim",
    [
        "资料没有提供安全认证信息，因此无法确认是否通过 ISO 27001 认证。",
        "简介未提供安全认证信息不等于未通过认证，不能据此推断认证状态。",
    ],
)
def test_unknown_certification_is_not_a_changed_measurement_or_reversed_fact(claim):
    quote = "此简介没有提供安全认证信息。"
    approved = CitationCheck(evidence_id="C0", statement=claim, status="supported")
    assert support_check("C0", claim, quote, approved).status == "supported"


@pytest.mark.parametrize(
    "formula,expected",
    [
        ("增长率按收入计算：(20 - 10) / 10 × 100 = 100%。", "supported"),
        ("2025 到 2026 年收入增长率：(20 − 10) ÷ 10 × 100% = 100%。", "supported"),
        ("增长率按收入计算：(20 - 10) / 10 × 100 = 200%。", "unsupported"),
        ("增加金额：20 万元 − 10 万元 = 20 万元。", "unsupported"),
        ("增长率按收入计算：(50 - 10) / 10 × 100 = 400%。", "unsupported"),
        ("收入增长率为 100%。", "unsupported"),
    ],
)
def test_derived_numbers_need_correct_arithmetic_and_cited_inputs(formula, expected):
    quote = "2025 年收入为 10 万元，2026 年收入为 20 万元。"
    approved = CitationCheck(evidence_id="C0", statement=formula, status="supported")
    assert support_check("C0", formula, quote, approved).status == expected


def test_report_keeps_verified_calculation_and_marks_it_as_inference(settings, knowledge_base, monkeypatch):
    runtime = build_runtime(
        task_id="growth", question="收入增长", settings=settings, knowledge_base=knowledge_base
    )
    runtime.llm.provider.name = "semantic-test"
    quote = "2025 年收入为 10 万元，2026 年收入为 20 万元。"
    calculation = "收入增长率：(20 - 10) / 10 * 100 = 100%。"
    evidence = Evidence(id="E1", claim=quote, quote=quote, source_id="local")
    report = FinalReport(
        title="收入",
        executive_summary=calculation,
        conclusions=[ReportClaim(statement=calculation, evidence_ids=["E1"])],
        sections=[ReportSection(heading="计算", body=calculation, evidence_ids=["E1"])],
    )
    monkeypatch.setattr(
        runtime.llm,
        "run",
        lambda *args, **kwargs: SimpleNamespace(
            value=VerificationReport(
                checks=[
                    CitationCheck(evidence_id=item["id"], statement=item["claim"], status="supported")
                    for item in kwargs["hints"]["evidence"]
                ]
            )
        ),
    )
    checked = WriterAgent(runtime)._audit_report(report, [evidence], _plan("收入增长"))
    assert checked.conclusions[0].statement == calculation
    assert checked.conclusions[0].assessment == "inference"
    assert checked.sections[0].body == calculation
    assert not checked.rejected_claims


def test_post_audit_deduplication_preserves_all_answers_in_summary(settings, knowledge_base):
    runtime = build_runtime(
        task_id="dedup", question="收入及上线日期", settings=settings, knowledge_base=knowledge_base
    )
    quotes = ["2025 年收入为 10 万元，2026 年收入为 20 万元。", "产品于 2026 年 9 月 10 日上线。"]
    evidence = [
        Evidence(id=f"E{index + 1}", claim=quote, quote=quote, source_id="local")
        for index, quote in enumerate(quotes)
    ]
    report = FinalReport(
        title="报告",
        executive_summary="无法核对的摘要",
        conclusions=[
            ReportClaim(statement=f"无法核对的第 {index} 条结论", evidence_ids=["E1", "E2"])
            for index in range(4)
        ],
    )
    checked = WriterAgent(runtime)._audit_report(report, evidence, _plan())
    assert len(checked.conclusions) == 2
    assert all(quote in checked.executive_summary for quote in quotes)


def test_references_merge_document_chunks_but_keep_distinct_quotes():
    from researchpilot.agents.writer import render_report
    from researchpilot.schemas import SourceRef

    sources = [
        SourceRef(id="chunk", kind="knowledge_base", title="年度记录", doc_id="document", chunk_id="c1"),
        SourceRef(id="full", kind="document", title="年度记录", doc_id="document"),
    ]
    evidence = [
        Evidence(id="E1", claim="收入 10 万元。", quote="收入 10 万元。", source_id="chunk"),
        Evidence(id="E2", claim="产品已上线。", quote="产品已上线。", source_id="full"),
    ]
    report = FinalReport(
        title="报告",
        executive_summary="",
        conclusions=[ReportClaim(statement=item.claim, evidence_ids=[item.id]) for item in evidence],
    )
    markdown = render_report(report, evidence, sources)
    assert markdown.count("1. **年度记录**") == 1
    assert "2. **年度记录**" not in markdown
    assert "| E1 | 收入 10 万元。 | [1] 年度记录 |" in markdown
    assert "| E2 | 产品已上线。 | [1] 年度记录 |" in markdown


def test_supported_income_summary_still_includes_date_used_only_in_analysis(settings, knowledge_base):
    runtime = build_runtime(
        task_id="summary-coverage",
        question="收入及上线日期",
        settings=settings,
        knowledge_base=knowledge_base,
    )
    income = "2025 年收入为 10 万元，2026 年收入为 20 万元。"
    date = "产品于 2026 年 9 月 10 日上线。"
    evidence = [
        Evidence(id="E1", claim=income, quote=income, source_id="local"),
        Evidence(id="E2", claim=date, quote=date, source_id="local"),
    ]
    report = FinalReport(
        title="报告",
        executive_summary=income + " [E1]",
        conclusions=[ReportClaim(statement=income, evidence_ids=["E1"])],
        sections=[ReportSection(heading="上线日期", body=date, evidence_ids=["E2"])],
    )
    checked = WriterAgent(runtime)._audit_report(report, evidence, _plan("收入及上线日期"))
    assert income in checked.executive_summary and date in checked.executive_summary
    assert "[E2]" in checked.executive_summary
    assert not checked.rejected_claims


def test_citation_aliases_merge_identical_document_quotes_without_merging_conflicts(settings, knowledge_base):
    from researchpilot.agents.writer import render_report

    runtime = build_runtime(
        task_id="aliases", question="离线部署", settings=settings, knowledge_base=knowledge_base
    )
    supported, denied = "本产品支持离线部署。", "本产品不支持离线部署。"
    for source, doc in [("chunk", "甲"), ("full", "甲"), ("other", "乙")]:
        runtime.sources.register_item(
            {"source_id": source, "doc_id": doc, "content": supported}, tool="document_reader"
        )
    evidence = [
        Evidence(id="E1", claim=supported, quote=supported, source_id="chunk"),
        Evidence(id="E2", claim=supported, quote=supported, source_id="full"),
        Evidence(id="E3", claim=supported, quote=supported, source_id="other"),
        Evidence(id="E4", claim=denied, quote=denied, source_id="full"),
    ]
    report = FinalReport(
        title="报告",
        executive_summary="本产品支持离线部署 [E1][E2]；另有资料 [E3]；不同原文 [E4]。",
        conclusions=[ReportClaim(statement=supported + " [E1, E2]", evidence_ids=["E1", "E2"])],
        sections=[ReportSection(heading="原文", body=denied + " [E4]", evidence_ids=["E3", "E4"])],
    )
    bound = WriterAgent(runtime)._bind_citations(report, evidence, VerificationReport(), set())
    assert bound.conclusions[0].evidence_ids == ["E1"]
    assert "[E2]" not in bound.executive_summary and "[E1] [E1]" not in bound.executive_summary
    assert bound.sections[0].evidence_ids == ["E3", "E4"]
    markdown = render_report(bound, evidence, runtime.sources.all())
    assert "| E2 |" not in markdown
    assert "| E1 |" in markdown and "| E3 |" in markdown and "| E4 |" in markdown


def test_extractive_fallback_does_not_truncate_last_answer(settings, knowledge_base):
    runtime = build_runtime(
        task_id="fallback", question="日期", settings=settings, knowledge_base=knowledge_base
    )
    quotes = [f"第 {index} 条记录。" for index in range(6)] + ["仅当批准时，于 2026 年 9 月 10 日上线。"]
    evidence = [
        Evidence(id=f"E{index + 1}", claim=quote, quote=quote, source_id="local")
        for index, quote in enumerate(quotes)
    ]
    report = WriterAgent(runtime)._fallback_report(_plan("日期"), evidence, VerificationReport())
    assert quotes[-1] in report.executive_summary
    assert report.conclusions[-1].evidence_ids == ["E7"]


def test_calculation_citations_do_not_hide_requested_year_amount(settings, knowledge_base, monkeypatch):
    runtime = build_runtime(
        task_id="year-mapping", question="收入与日期", settings=settings, knowledge_base=knowledge_base
    )
    runtime.llm.provider.name = "semantic-test"
    quotes = ["2025 年收入为 10 万元。", "2026 年收入为 20 万元。", "产品于 2026 年 9 月 10 日上线。"]
    evidence = [
        Evidence(id=f"E{index + 1}", claim=quote, quote=quote, source_id="local")
        for index, quote in enumerate(quotes)
    ]
    calculation = "增长率：(20 - 10) / 10 * 100 = 100%。"
    report = FinalReport(
        title="收入",
        executive_summary=quotes[0] + quotes[2] + calculation,
        conclusions=[
            ReportClaim(statement=quotes[0], evidence_ids=["E1"]),
            ReportClaim(statement=quotes[2], evidence_ids=["E3"]),
            ReportClaim(statement=calculation, evidence_ids=["E1", "E2"]),
        ],
        sections=[ReportSection(heading="原始收入", body="".join(quotes), evidence_ids=["E1", "E2", "E3"])],
    )
    monkeypatch.setattr(
        runtime.llm,
        "run",
        lambda *a, **kw: SimpleNamespace(
            value=VerificationReport(
                checks=[
                    CitationCheck(evidence_id=item["id"], statement=item["claim"], status="supported")
                    for item in kw["hints"]["evidence"]
                ]
            )
        ),
    )
    checked = WriterAgent(runtime)._audit_report(report, evidence, _plan("收入与日期"))
    assert all(quote in checked.executive_summary for quote in quotes)
    assert calculation in checked.executive_summary


def test_evaluation_does_not_call_an_existing_but_unsupported_citation_correct():
    from researchpilot.evaluation.dataset import GoldenTask
    from researchpilot.evaluation.judge import judge_task
    from researchpilot.schemas import ResearchResult, SourceRef

    result = ResearchResult(
        task_id="citation-regression",
        trace_id="citation-regression-trace",
        question="离线部署",
        status="completed",
        report=FinalReport(title="报告", executive_summary="支持离线部署", markdown="支持离线部署 [E1]"),
        evidence=EvidenceBundle(
            evidence=[Evidence(id="E1", claim="支持离线部署", quote="不支持离线部署", source_id="local")],
            sources=[SourceRef(id="local", title="资料", kind="document")],
        ),
        verification=VerificationReport(
            checks=[CitationCheck(evidence_id="E1", statement="支持离线部署", status="unsupported")]
        ),
    )
    judged = judge_task(
        GoldenTask(id="trust", category="citation", question="离线部署"), result, None, latency_s=0
    )
    assert judged.citation_integrity == 1
    assert judged.citation_correctness == 0 and not judged.passed


def test_followup_retrieval_runs_only_the_requested_gap(settings, knowledge_base, monkeypatch):
    runtime = build_runtime(task_id="gap", question="补检", settings=settings, knowledge_base=knowledge_base)
    agent = ResearchAgent(runtime)
    asked: list[str] = []

    def collect(subtask, ctx):
        asked.append(subtask.question)
        return []

    monkeypatch.setattr(agent, "_execute_subtask", collect)
    agent.run(_plan(), follow_up_queries=["仅补检价格信息"], iteration=2)
    assert asked == ["仅补检价格信息"]


def test_selected_scope_and_followup_context_reach_the_model(settings, knowledge_base, monkeypatch):
    ids = [knowledge_base.summaries()[0].doc_id]
    provider = MockLLMProvider(settings)
    pipeline = ResearchPipeline(settings, provider=provider, knowledge_base=knowledge_base)
    stages: list[str] = []
    publish = pipeline._publish_stage

    def record_stage(task, stage):
        stages.append(stage)
        publish(task, stage)

    monkeypatch.setattr(pipeline, "_publish_stage", record_stage)
    captured: list[str] = []
    complete = provider.complete

    def capture(messages, **kwargs):
        captured.extend(message.content for message in messages if "previous-research" in message.content)
        return complete(messages, **kwargs)

    monkeypatch.setattr(provider, "complete", capture)
    try:
        result = pipeline.run(ResearchRequest(question="资料中的工具和协议是什么？", document_ids=ids))
        assert result.status == "completed" and result.document_ids == ids
        assert all(source.doc_id in ids for source in result.evidence.sources if source.doc_id)
        assert result.metrics.mcp_calls == 0
        assert "planning" in stages and "checking_report" in stages
        assert result.report and "confidence=" not in result.report.markdown
        followup = pipeline.run(
            ResearchRequest(question="这些资料还有什么风险？", parent_task_id=result.task_id)
        )
        assert followup.parent_task_id == result.task_id and followup.document_ids == ids
        assert captured and result.question in captured[0]
        assert "quote" in captured[0] and "source_id" in captured[0]
        with pytest.raises(ResearchScopeError):
            pipeline.run(
                ResearchRequest(
                    question="更换资料范围进行追问", parent_task_id=result.task_id, document_ids=[]
                )
            )
    finally:
        pipeline.close()


def test_snapshot_prevents_document_reader_and_retrieval_from_seeing_unselected_docs(knowledge_base):
    ids = [knowledge_base.summaries()[0].doc_id]
    scoped = knowledge_base.scoped(ids)
    unselected = next(doc for doc in knowledge_base.summaries() if doc.doc_id not in ids)
    assert scoped.get_document(unselected.doc_id) is None
    assert all(hit.chunk.doc_id in ids for hit in scoped.search("知识库工具和协议", rewrite=False).hits)


def test_docx_import_keeps_paragraphs_and_tables_in_order():
    from docx import Document

    document = Document()
    document.add_paragraph("第一段：个人研究资料。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "价格", "10 元"
    document.add_paragraph("最后一段：不支持离线。")
    data = BytesIO()
    document.save(data)
    loaded = DocumentLoader().load_bytes(data.getvalue(), filename="说明.docx")[0].content
    assert loaded.index("第一段") < loaded.index("10 元") < loaded.index("最后一段")


def text_pdf(text="Offline deployment is not supported."):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 50 750 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    data = BytesIO()
    writer.write(data)
    return data.getvalue()


def test_pdf_import_keeps_page_marker_and_source_text():
    loaded = DocumentLoader().load_bytes(text_pdf(), filename="spec.pdf")[0]
    assert "第 1 页" in loaded.content and "not supported" in loaded.content
    assert loaded.title == "spec"


def test_empty_scanned_or_corrupt_pdf_explains_the_limit():
    from pypdf import PdfWriter

    writer, data = PdfWriter(), BytesIO()
    writer.add_blank_page(width=100, height=100)
    writer.write(data)
    with pytest.raises(ValueError, match="OCR"):
        DocumentLoader().load_bytes(data.getvalue(), filename="scan.pdf")
    with pytest.raises(ValueError, match="PDF"):
        DocumentLoader().load_bytes(b"broken", filename="broken.pdf")
